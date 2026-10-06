#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
csv_a_excel.py
===============
Convierte archivos CSV a Excel (.xlsx) con formato profesional y limpieza
automática de "datos basura" (columnas vacías, filas vacías, celdas con
espacios sueltos, texto "True"/"False" convertido a booleano real, etc).

Pensado para CSV de resultados tipo "una fila por objeto analizado, muchas
columnas de distinto tipo" (como resultados_exoplanetas.csv), pero funciona
con cualquier CSV.

Qué hace por vos:
  - Detecta el separador y la codificación del CSV automáticamente.
  - Elimina columnas y filas 100% vacías.
  - Limpia espacios sobrantes/dobles en texto.
  - Convierte "True"/"False" a booleano real (con color: verde/rojo).
  - Detecta columnas numéricas, de fecha, de porcentaje, y las formatea.
  - Agrupa columnas por tema (identificación, posición, observación,
    detección, resultado, etc.) y les pone un color de encabezado distinto,
    para que se entienda de un vistazo qué es cada bloque de datos.
  - Ajusta el ancho de columnas, congela el encabezado y agrega autofiltro.

Requisitos (una sola vez):
    pip install pandas openpyxl

Funciona igual en CachyOS, Linux Mint (o cualquier Linux) y macOS: es un
script de Python puro, no usa nada específico del sistema operativo.

Uso:
    python3 csv_a_excel.py archivo.csv
    python3 csv_a_excel.py archivo.csv -o resultado.xlsx
    python3 csv_a_excel.py *.csv --out-dir convertidos/
    python3 csv_a_excel.py archivo.csv --umbral-relleno 0.3
    python3 csv_a_excel.py archivo.csv --quitar-duplicados
    python3 csv_a_excel.py archivo.csv --sin-limpieza
    python3 csv_a_excel.py archivo.csv --sin-color
"""

import argparse
import csv
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Grupos de columnas: (nombre del grupo, color de encabezado, palabras clave)
# Si el nombre de una columna contiene alguna de las palabras clave (sin
# importar mayúsculas/acentos), se le asigna el color de ese grupo. Si no
# matchea ninguno, usa el grupo "General". Ajustá esta lista si tus CSV
# tienen otras columnas típicas.
# --------------------------------------------------------------------------
GRUPOS_COLUMNAS = [
    ("Identificación", "1F4E78", ["id", "nombre", "name", "codigo", "code"]),
    ("Posición", "0F6B72", ["ra_", "dec_", "lat", "lon", "coord", "posicion"]),
    ("Parámetros físicos", "B45309", ["teff", "radio", "masa", "mag", "densidad",
                                       "temperatura", "tmag", "peso", "volumen"]),
    ("Observación", "6B3FA0", ["sector", "observ", "fecha", "analisis",
                                "curva", "disponible", "puntos"]),
    ("Detección / Señal", "2E7D32", ["periodo", "snr", "profundidad",
                                      "confianza", "anomalia", "transito",
                                      "binaria", "senal", "señal"]),
    ("Catálogo / Confirmación", "9A7B00", ["confirmado", "conocid",
                                            "publicad", "error"]),
    ("Resultado", "8B2E2E", ["estado", "mensaje", "status", "resultado",
                              "observacion_final", "comentario"]),
]
COLOR_GENERAL = "44546A"  # gris azulado, para columnas que no matchean nada

FUENTE = "Arial"


def normalizar(texto):
    """Minúsculas y sin tildes, para comparar nombres de columnas."""
    reemplazos = str.maketrans("áéíóúñ", "aeioun")
    return texto.lower().translate(reemplazos)


def grupo_de_columna(nombre_columna):
    n = normalizar(nombre_columna)
    partes = re.split(r"[^a-z0-9]+", n)  # tokens separados por "_"
    for nombre_grupo, color, claves in GRUPOS_COLUMNAS:
        for clave in claves:
            if clave.endswith("_"):
                # prefijos tipo "ra_", "dec_": deben estar al inicio
                if n.startswith(clave):
                    return nombre_grupo, color
            elif len(clave) <= 3:
                # palabras muy cortas ("id") solo matchean como token
                # completo, para no confundir "profundidad" con "id"
                if clave in partes:
                    return nombre_grupo, color
            elif clave in n:
                # palabras más largas: substring alcanza (cubre plurales,
                # ej. "disponible" en "disponibles")
                return nombre_grupo, color
    return "General", COLOR_GENERAL


# --------------------------------------------------------------------------
# Lectura robusta del CSV: detecta separador y codificación.
# --------------------------------------------------------------------------
def leer_csv(ruta):
    ruta = Path(ruta)
    # Probamos utf-8-sig primero (soporta el BOM que agregan Excel/Windows),
    # y si falla, latin-1 como red de seguridad.
    for codificacion in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(ruta, "r", encoding=codificacion, newline="") as f:
                muestra = f.read(8192)
            break
        except UnicodeDecodeError:
            continue
    else:
        codificacion = "latin-1"

    try:
        separador = csv.Sniffer().sniff(muestra, delimiters=",;\t|").delimiter
    except csv.Error:
        separador = ","

    df = pd.read_csv(ruta, sep=separador, encoding=codificacion,
                      dtype=str, keep_default_na=True, na_values=[
                          "", "NA", "N/A", "NULL", "null", "None", "none",
                          "-", "--", "nan", "NaN",
                      ])
    # Nombres de columna sin espacios sobrantes ni BOM residual
    df.columns = [str(c).strip().lstrip("\ufeff") for c in df.columns]
    return df


# --------------------------------------------------------------------------
# Limpieza de datos
# --------------------------------------------------------------------------
def limpiar_dataframe(df, umbral_relleno=None, quitar_duplicados=False,
                       silencioso=False):
    reporte = []
    filas_iniciales, columnas_iniciales = df.shape

    # 1) Limpiar espacios en texto ("TYC  935-859-1" -> "TYC 935-859-1")
    for col in df.columns:
        if pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object:
            df[col] = df[col].apply(
                lambda v: re.sub(r"\s+", " ", v.strip()) if isinstance(v, str) else v
            )
            df[col] = df[col].replace("", pd.NA)

    # 2) Columnas 100% vacías -> afuera (son "basura": no aportan nada)
    cols_vacias = [c for c in df.columns if df[c].isna().all()]
    if cols_vacias:
        df = df.drop(columns=cols_vacias)
        reporte.append(f"Columnas vacías eliminadas: {', '.join(cols_vacias)}")

    # 3) Filas 100% vacías -> afuera
    filas_vacias = df.isna().all(axis=1).sum()
    if filas_vacias:
        df = df.dropna(how="all")
        reporte.append(f"Filas completamente vacías eliminadas: {filas_vacias}")

    # 4) Duplicados exactos (opcional, porque a veces son datos válidos repetidos)
    if quitar_duplicados:
        antes = len(df)
        df = df.drop_duplicates()
        eliminadas = antes - len(df)
        if eliminadas:
            reporte.append(f"Filas duplicadas eliminadas: {eliminadas}")

    # 5) Filas con muy poco contenido útil (opcional, vía --umbral-relleno)
    if umbral_relleno is not None:
        relleno = df.notna().sum(axis=1) / max(len(df.columns), 1)
        antes = len(df)
        df = df[relleno >= umbral_relleno]
        eliminadas = antes - len(df)
        if eliminadas:
            reporte.append(
                f"Filas con menos de {umbral_relleno:.0%} de datos completos "
                f"eliminadas: {eliminadas}"
            )

    # 6) Booleanos reales en vez de texto "True"/"False"
    mapa_bool = {"true": True, "false": False, "verdadero": True, "falso": False,
                 "si": True, "no": False, "sí": True}
    for col in df.columns:
        if not (pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object):
            continue
        valores = df[col].dropna().unique()
        if len(valores) == 0:
            continue
        if all(str(v).strip().lower() in mapa_bool for v in valores):
            df[col] = df[col].apply(
                lambda v: mapa_bool[str(v).strip().lower()] if pd.notna(v) else v
            )

    # 7) Columnas numéricas: si la gran mayoría de los valores no vacíos
    #    convierten a número, tratamos la columna como numérica.
    for col in df.columns:
        if not (pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object):
            continue
        no_nulos = df[col].dropna()
        if len(no_nulos) == 0:
            continue
        convertidos = pd.to_numeric(no_nulos, errors="coerce")
        tasa_exito = convertidos.notna().mean()
        if tasa_exito >= 0.95:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    # 8) Columnas de fecha/hora (formato ISO típico de exportaciones)
    for col in df.columns:
        if not (pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object):
            continue
        n = normalizar(col)
        if not any(p in n for p in ["fecha", "date", "time", "inicio", "fin"]):
            continue
        no_nulos = df[col].dropna()
        if len(no_nulos) == 0:
            continue
        convertidos = pd.to_datetime(no_nulos, errors="coerce", utc=False,
                                      format="mixed")
        if convertidos.notna().mean() >= 0.9:
            serie_fecha = pd.to_datetime(df[col], errors="coerce", format="mixed")
            if getattr(serie_fecha.dt, "tz", None) is not None:
                # Excel no admite fechas con huso horario: lo quitamos
                # (la hora queda igual, solo se descarta el offset UTC).
                serie_fecha = serie_fecha.dt.tz_localize(None)
            df[col] = serie_fecha

    filas_finales, columnas_finales = df.shape
    if not silencioso:
        print(f"  Filas: {filas_iniciales} -> {filas_finales}  |  "
              f"Columnas: {columnas_iniciales} -> {columnas_finales}")
        for linea in reporte:
            print(f"    - {linea}")

    return df


# --------------------------------------------------------------------------
# Formato numérico según el tipo/nombre de columna
# --------------------------------------------------------------------------
def formato_numero(nombre_col, serie):
    n = normalizar(nombre_col)
    if "pct" in n or "porcentaje" in n or "%" in nombre_col:
        return '0.00"%"'

    no_nulos = serie.dropna()
    es_entera = pd.api.types.is_integer_dtype(serie) or (
        len(no_nulos) > 0 and (no_nulos % 1 == 0).all()
    )
    if es_entera:
        return "#,##0"

    maximo = no_nulos.abs().max() if len(no_nulos) else None
    if maximo is None or pd.isna(maximo):
        return "0.00"
    if maximo < 1:
        return "0.000000"
    if maximo < 100:
        return "0.0000"
    return "#,##0.00"


# --------------------------------------------------------------------------
# Construcción del Excel formateado
# --------------------------------------------------------------------------
def escribir_excel(df, ruta_salida, nombre_hoja="Datos", usar_color=True):
    wb = Workbook()
    ws = wb.active
    ws.title = nombre_hoja[:31] or "Datos"

    columnas = list(df.columns)
    grupos = [grupo_de_columna(c) if usar_color else ("General", COLOR_GENERAL)
              for c in columnas]

    borde_fino = Side(style="thin", color="D9D9D9")
    borde = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)

    # --- Encabezado ---
    for j, (col, (_, color)) in enumerate(zip(columnas, grupos), start=1):
        celda = ws.cell(row=1, column=j, value=col)
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=color)
        celda.alignment = Alignment(horizontal="center", vertical="center",
                                     wrap_text=True)
        celda.border = borde
    ws.row_dimensions[1].height = 30

    # --- Datos ---
    formatos = {}
    for col in columnas:
        if pd.api.types.is_numeric_dtype(df[col]):
            formatos[col] = formato_numero(col, df[col])
        elif pd.api.types.is_datetime64_any_dtype(df[col]):
            formatos[col] = "yyyy-mm-dd hh:mm"

    fill_par = PatternFill("solid", fgColor="F2F2F2")
    fill_bool_true = PatternFill("solid", fgColor="C6EFCE")
    fill_bool_false = PatternFill("solid", fgColor="FFC7CE")

    for i, (_, fila) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(columnas, start=1):
            valor = fila[col]
            if pd.isna(valor):
                valor = None
            elif isinstance(valor, bool):
                pass  # openpyxl entiende bool nativo
            elif hasattr(valor, "to_pydatetime"):
                valor = valor.to_pydatetime()

            celda = ws.cell(row=i, column=j, value=valor)
            celda.font = Font(name=FUENTE, size=10)
            celda.border = borde
            celda.alignment = Alignment(vertical="center",
                                         wrap_text=(col.lower() in
                                                    ("mensaje", "comentario",
                                                     "descripcion", "descripción")))

            if col in formatos:
                celda.number_format = formatos[col]

            if isinstance(valor, bool):
                celda.fill = fill_bool_true if valor else fill_bool_false
                celda.alignment = Alignment(horizontal="center", vertical="center")
            elif i % 2 == 0:
                celda.fill = fill_par

    # --- Ancho de columnas ---
    for j, col in enumerate(columnas, start=1):
        letra = get_column_letter(j)
        largo_header = len(str(col))
        try:
            largo_datos = int(df[col].astype(str).str.len().quantile(0.95))
        except Exception:
            largo_datos = 10
        ancho = max(largo_header, largo_datos) + 2
        ws.column_dimensions[letra].width = min(max(ancho, 10), 45)

    # --- Encabezado fijo + autofiltro ---
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)
    return ruta_salida


# --------------------------------------------------------------------------
# Programa principal
# --------------------------------------------------------------------------
def procesar_archivo(ruta_csv, args):
    print(f"\n{ruta_csv}")
    df = leer_csv(ruta_csv)

    if not args.sin_limpieza:
        df = limpiar_dataframe(
            df,
            umbral_relleno=args.umbral_relleno,
            quitar_duplicados=args.quitar_duplicados,
        )
    else:
        print(f"  (sin limpieza) Filas: {len(df)}  Columnas: {len(df.columns)}")

    if args.out and len(args.csv) == 1:
        salida = Path(args.out)
    else:
        carpeta = Path(args.out_dir) if args.out_dir else Path(ruta_csv).parent
        salida = carpeta / (Path(ruta_csv).stem + ".xlsx")

    ruta_final = escribir_excel(df, salida, nombre_hoja=args.hoja,
                                 usar_color=not args.sin_color)
    print(f"  -> {ruta_final}")
    return ruta_final


def main():
    parser = argparse.ArgumentParser(
        description="Convierte CSV a Excel formateado, limpiando datos basura.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("csv", nargs="+", help="Archivo(s) CSV a convertir (admite comodines *.csv)")
    parser.add_argument("-o", "--out", help="Ruta del archivo .xlsx de salida (solo si es un único CSV)")
    parser.add_argument("--out-dir", help="Carpeta de salida para varios archivos")
    parser.add_argument("--hoja", default="Datos", help="Nombre de la hoja de Excel (default: Datos)")
    parser.add_argument("--sin-limpieza", action="store_true",
                         help="No limpiar nada: solo convertir y formatear tal cual está el CSV")
    parser.add_argument("--umbral-relleno", type=float, default=None,
                         help="Elimina filas con menos de esta fracción (0-1) de celdas con datos. "
                              "Ej: 0.3 descarta filas con menos del 30%% de columnas completas")
    parser.add_argument("--quitar-duplicados", action="store_true",
                         help="Elimina filas exactamente duplicadas")
    parser.add_argument("--sin-color", action="store_true",
                         help="No colorear encabezados por grupo temático")
    args = parser.parse_args()

    for patron in args.csv:
        rutas = sorted(Path().glob(patron)) if any(c in patron for c in "*?[]") else [Path(patron)]
        if not rutas:
            print(f"No se encontró: {patron}", file=sys.stderr)
            continue
        for ruta in rutas:
            if not ruta.exists():
                print(f"No se encontró: {ruta}", file=sys.stderr)
                continue
            try:
                procesar_archivo(ruta, args)
            except Exception as e:
                print(f"  ERROR procesando {ruta}: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
